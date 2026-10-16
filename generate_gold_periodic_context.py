"""Inspect official World Gold Council periodic XLSX downloads; dry-run only."""
import io
import re
import urllib.error
import urllib.request
from datetime import datetime, timezone
from html.parser import HTMLParser
from urllib.parse import urljoin, urlparse

from openpyxl import load_workbook

ETF_PAGE = "https://www.gold.org/goldhub/data/gold-etfs-holdings-and-flows"
RESERVES_PAGE = "https://www.gold.org/goldhub/data/gold-reserves-by-country"
OFFICIAL_HOSTS = {"www.gold.org", "gold.org"}
TIMEOUT = 25
MAX_DOWNLOAD_BYTES = 25 * 1024 * 1024


class PageParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.title = []
        self.text = []
        self.links = []
        self._in_title = False
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "title":
            self._in_title = True
        if tag in {"script", "style", "noscript", "svg"}:
            self._skip_depth += 1
        if tag == "a" and attrs.get("href"):
            self.links.append({"href": attrs["href"], "label": ""})

    def handle_endtag(self, tag):
        if tag == "title":
            self._in_title = False
        if tag in {"script", "style", "noscript", "svg"} and self._skip_depth:
            self._skip_depth -= 1

    def handle_data(self, data):
        value = " ".join(data.split())
        if not value or self._skip_depth:
            return
        if self._in_title:
            self.title.append(value)
        self.text.append(value)
        if self.links and not self.links[-1]["label"]:
            self.links[-1]["label"] = value


def request(url):
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname not in OFFICIAL_HOSTS:
        raise ValueError("Blocked non-official or non-HTTPS source URL")
    req = urllib.request.Request(url, headers={"User-Agent": "FundBoardGoldBot/1.0 (research draft)"})
    return urllib.request.urlopen(req, timeout=TIMEOUT)


def fetch_page(url):
    with request(url) as response:
        final_url = response.geturl()
        if urlparse(final_url).scheme != "https" or urlparse(final_url).hostname not in OFFICIAL_HOSTS:
            raise ValueError("Unexpected final URL for World Gold Council source")
        raw = response.read(MAX_DOWNLOAD_BYTES + 1)
        if len(raw) > MAX_DOWNLOAD_BYTES:
            raise ValueError("Source page exceeds configured size limit")
    parser = PageParser()
    parser.feed(raw.decode("utf-8", errors="replace"))
    return {
        "url": final_url,
        "title": " ".join(parser.title).strip() or final_url,
        "text": " ".join(" ".join(parser.text).split()),
        "links": [
            {"url": urljoin(final_url, link["href"]), "label": link["label"]}
            for link in parser.links
            if urlparse(urljoin(final_url, link["href"])).hostname in OFFICIAL_HOSTS
            and urlparse(urljoin(final_url, link["href"])).scheme == "https"
        ],
    }


def find_xlsx_links(page, keywords):
    found = []
    for link in page["links"]:
        target = link["url"].lower().split("?", 1)[0]
        label = link["label"].lower()
        parsed = urlparse(link["url"])
        if parsed.scheme != "https" or parsed.hostname not in OFFICIAL_HOSTS:
            continue
        if target.endswith(".xlsx") or ".xlsx/" in target or "xlsx" in target or "xlsx" in label:
            if any(word in label or word in target for word in keywords):
                found.append(link)
    # Keep distinct links in page order.
    unique = []
    seen = set()
    for link in found:
        if link["url"] not in seen:
            seen.add(link["url"])
            unique.append(link)
    return unique


def inspect_workbook(url):
    try:
        with request(url) as response:
            final_url = response.geturl()
            if urlparse(final_url).scheme != "https" or urlparse(final_url).hostname not in OFFICIAL_HOSTS:
                raise ValueError("Workbook redirected away from official host")
            raw = response.read(MAX_DOWNLOAD_BYTES + 1)
            if len(raw) > MAX_DOWNLOAD_BYTES:
                return {"url": url, "status": "too_large"}
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheets = []
        for sheet in workbook.worksheets[:8]:
            iterator = sheet.iter_rows(values_only=True)
            headers = next(iterator, ())
            sample = []
            for row in iterator:
                if any(value is not None for value in row):
                    sample.append(row)
                if len(sample) >= 3:
                    break
            sheets.append({
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "headers": [str(v)[:100] if v is not None else "" for v in headers[:12]],
                "sample_rows": [[str(v)[:100] if v is not None else "" for v in row[:12]] for row in sample],
            })
        workbook.close()
        return {"url": url, "status": "parsed", "bytes": len(raw), "sheets": sheets}
    except urllib.error.HTTPError as exc:
        if exc.code in {401, 403}:
            return {"url": url, "status": f"access_restricted_http_{exc.code}"}
        return {"url": url, "status": f"http_error_{exc.code}"}
    except Exception as exc:
        return {"url": url, "status": f"unavailable_{type(exc).__name__}"}


def format_workbook(label, result):
    lines = [f"- {label}: {result['url']}", f"  وضعیت دریافت/پردازش: {result['status']}"]
    if result["status"] == "parsed":
        lines.append(f"  حجم فایل: {result['bytes']} بایت")
        for sheet in result["sheets"]:
            lines.append(f"  برگه «{sheet['name']}»: حدود {sheet['max_row']} ردیف، {sheet['max_column']} ستون")
            if sheet["headers"]:
                lines.append("  سرستون‌ها: " + " | ".join(sheet["headers"]))
            for row in sheet["sample_rows"]:
                lines.append("  نمونه: " + " | ".join(row))
    elif result["status"].startswith("access_restricted"):
        lines.append("  فایل از این محیط بدون دسترسی مجاز قابل دریافت نیست؛ هیچ عددی از آن استخراج نشده است.")
    else:
        lines.append("  داده عددی استخراج نشد؛ دسترسی یا ساختار فایل نیازمند بررسی است.")
    return lines


def build_report(etf_page, reserves_page, etf_files, reserve_files, generated_at=None):
    generated_at = generated_at or datetime.now(timezone.utc)
    lines = [
        "🟡 فاندبورد طلا | بررسی فایل‌های داده دوره‌ای شورای جهانی طلا",
        f"زمان ساخت (UTC): {generated_at.astimezone(timezone.utc).isoformat()}",
        "این گزارش بررسی فنی فایل‌هاست؛ ارقام نمونه یا سرستون‌ها به‌عنوان تغییرات قطعی بازار تفسیر نشده‌اند.",
        "",
        "📊 ETFهای طلای جهانی",
        f"صفحه رسمی: {etf_page['url']}",
        f"عنوان صفحه: {etf_page['title']}",
        "طبق روش انتشار منبع، داده‌های وب می‌توانند هفتگی و ماهانه باشند و فایل Excel معمولاً ماهانه منتشر می‌شود.",
    ]
    if etf_files:
        for item in etf_files[:3]:
            lines.extend(format_workbook("فایل داده ETF", item))
    else:
        lines.append("لینک فایل Excel از صفحه به‌صورت مطمئن شناسایی نشد؛ ارقام جریان سرمایه گزارش نمی‌شود.")
    lines.extend([
        "",
        "🏦 ذخایر طلای بانک‌های مرکزی",
        f"صفحه رسمی: {reserves_page['url']}",
        f"عنوان صفحه: {reserves_page['title']}",
        "داده‌های ذخایر و تغییرات رسمی ممکن است ماهانه یا فصلی و با تأخیر منتشر شوند.",
    ])
    if reserve_files:
        for item in reserve_files[:3]:
            lines.extend(format_workbook("فایل داده ذخایر", item))
    else:
        lines.append("لینک فایل Excel از صفحه به‌صورت مطمئن شناسایی نشد؛ تغییر خریدوفروش محاسبه نشده است.")
    lines.extend([
        "",
        "مرحله بعدی پس از تأیید ساختار و دسترسی فایل‌ها: تعریف ستون‌های معتبر، تشخیص تاریخ مرجع، کنترل واحدها و محاسبه تغییرات فقط بین دوره‌های قابل مقایسه.",
        "⚠️ محتوای آموزشی و اطلاعاتی؛ نه توصیه خرید یا فروش.",
        "Dry-run only: no Telegram message was sent.",
    ])
    return "\n".join(lines)


def main():
    warnings = []
    try:
        etf_page = fetch_page(ETF_PAGE)
    except Exception as exc:
        warnings.append(f"ETF page unavailable: {type(exc).__name__}: {exc}")
        etf_page = {"url": ETF_PAGE, "title": "Unavailable", "text": "", "links": []}
    try:
        reserves_page = fetch_page(RESERVES_PAGE)
    except Exception as exc:
        warnings.append(f"Reserves page unavailable: {type(exc).__name__}: {exc}")
        reserves_page = {"url": RESERVES_PAGE, "title": "Unavailable", "text": "", "links": []}

    etf_links = find_xlsx_links(etf_page, ["etf", "flow", "holdings"]) if etf_page["text"] else []
    reserve_links = find_xlsx_links(reserves_page, ["reserve", "gold", "change", "official"]) if reserves_page["text"] else []
    etf_results = [inspect_workbook(link["url"]) for link in etf_links[:3]]
    reserve_results = [inspect_workbook(link["url"]) for link in reserve_links[:3]]
    for warning in warnings:
        print("WARNING:", warning)
    report = build_report(etf_page, reserves_page, etf_results, reserve_results)
    with open("gold_periodic_context_draft.txt", "w", encoding="utf-8") as output:
        output.write(report + "\n")
    print(report)


if __name__ == "__main__":
    main()
